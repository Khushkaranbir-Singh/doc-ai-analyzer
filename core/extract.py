"""Read text and metadata out of uploaded files.

Every extractor degrades gracefully: if an optional dependency or an OCR engine
is missing, the file still loads and the UI explains what is unavailable.
"""

from __future__ import annotations

import hashlib
import io
import os
import re
from dataclasses import dataclass, field
from typing import List

SUPPORTED_EXTENSIONS = ["pdf", "doc", "docx", "txt", "md", "jpg", "jpeg", "png", "exe"]

TEXT_EXTENSIONS = {"txt", "md", "csv", "log"}
IMAGE_EXTENSIONS = {"jpg", "jpeg", "png"}
BINARY_EXTENSIONS = {"exe", "dll", "bin"}

MAX_FILE_MB = 100


@dataclass
class Document:
    """Everything the rest of the app needs to know about one upload."""

    name: str
    extension: str
    size_bytes: int
    sha256: str
    text: str = ""
    pages: List[str] = field(default_factory=list)
    kind: str = "document"          # document | image | binary
    notes: List[str] = field(default_factory=list)
    meta: dict = field(default_factory=dict)

    @property
    def size_mb(self) -> float:
        return self.size_bytes / (1024 * 1024)

    @property
    def has_text(self) -> bool:
        return len(self.text.strip()) >= 40

    @property
    def page_count(self) -> int:
        return len(self.pages)


def _sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _clean(text: str) -> str:
    text = text.replace("\x00", " ").replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t\u00a0]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    # Repair words broken across a line end: "analy-\nsis" -> "analysis"
    text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)
    return text.strip()


# --------------------------------------------------------------------------- #
# Per-format readers
# --------------------------------------------------------------------------- #

def _read_pdf(data: bytes, doc: Document) -> None:
    pages: List[str] = []
    try:
        import pdfplumber

        with pdfplumber.open(io.BytesIO(data)) as pdf:
            doc.meta.update({k: str(v) for k, v in (pdf.metadata or {}).items() if v})
            for page in pdf.pages:
                pages.append(page.extract_text() or "")
    except Exception as exc:  # pdfplumber missing or file is unusual
        doc.notes.append(f"Detailed PDF reader unavailable ({exc}). Using the basic reader.")
        try:
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(data))
            if reader.metadata:
                doc.meta.update(
                    {str(k).lstrip("/"): str(v) for k, v in reader.metadata.items() if v}
                )
            pages = [(p.extract_text() or "") for p in reader.pages]
        except Exception as exc2:
            doc.notes.append(f"This PDF could not be read: {exc2}")
            pages = []

    doc.pages = [_clean(p) for p in pages]
    doc.text = _clean("\n\n".join(doc.pages))
    if pages and not doc.has_text:
        doc.notes.append(
            "No selectable text found. This looks like a scanned PDF — "
            "export it with OCR, or upload the pages as images instead."
        )


def _read_docx(data: bytes, doc: Document) -> None:
    try:
        import docx  # python-docx

        d = docx.Document(io.BytesIO(data))
        blocks = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells]
                if any(cells):
                    blocks.append(" | ".join(cells))
        props = d.core_properties
        for key in ("title", "author", "subject", "created", "modified"):
            value = getattr(props, key, None)
            if value:
                doc.meta[key] = str(value)
        doc.text = _clean("\n".join(blocks))
        doc.pages = _split_into_pages(doc.text)
    except Exception as exc:
        doc.notes.append(f"This Word file could not be read: {exc}")


def _read_legacy_doc(data: bytes, doc: Document) -> None:
    """Best-effort reader for the old binary .doc format."""
    doc.notes.append(
        "Legacy .doc files store text in a binary format, so extraction is "
        "approximate. Save the file as .docx or .pdf for a clean read."
    )
    raw = data.decode("latin-1", errors="ignore")
    chunks = re.findall(r"[ -~\n]{6,}", raw)
    text = "\n".join(c for c in chunks if re.search(r"[A-Za-z]{3,}", c))
    text = re.sub(r"(HYPERLINK|PAGEREF|MERGEFORMAT|Microsoft Word|Normal\.dotm?)[^\n]*", " ", text)
    doc.text = _clean(text)
    doc.pages = _split_into_pages(doc.text)


def _read_plain(data: bytes, doc: Document) -> None:
    for encoding in ("utf-8", "utf-16", "latin-1"):
        try:
            doc.text = _clean(data.decode(encoding))
            break
        except UnicodeDecodeError:
            continue
    doc.pages = _split_into_pages(doc.text)


def _read_image(data: bytes, doc: Document) -> None:
    doc.kind = "image"
    try:
        from PIL import Image

        image = Image.open(io.BytesIO(data))
        doc.meta.update(
            {
                "format": image.format or "unknown",
                "mode": image.mode,
                "width": image.width,
                "height": image.height,
                "megapixels": round(image.width * image.height / 1_000_000, 2),
            }
        )
    except Exception as exc:
        doc.notes.append(f"This image could not be opened: {exc}")
        return

    try:
        import pytesseract

        text = pytesseract.image_to_string(image)
        doc.text = _clean(text)
        if doc.has_text:
            doc.notes.append("Text was read from this image with OCR.")
        else:
            doc.notes.append("OCR ran but found no readable text in this image.")
    except Exception:
        doc.notes.append(
            "OCR is not installed here, so the image loaded without text. "
            "Add tesseract-ocr to packages.txt to read text from images."
        )
    doc.pages = _split_into_pages(doc.text) if doc.text else []


def _read_binary(data: bytes, doc: Document) -> None:
    """Inspect an executable without running it. Nothing here executes code."""
    doc.kind = "binary"
    doc.meta["signature"] = data[:2].decode("latin-1", errors="ignore")
    doc.meta["is_windows_pe"] = data[:2] == b"MZ"
    doc.meta["entropy"] = round(_entropy(data[: 2_000_000]), 3)
    strings = re.findall(rb"[\x20-\x7e]{6,}", data[: 5_000_000])
    readable = [s.decode("latin-1") for s in strings]
    doc.meta["printable_strings"] = len(readable)
    doc.text = _clean("\n".join(readable[:4000]))
    doc.pages = _split_into_pages(doc.text)
    doc.notes.append(
        "Executables are inspected, never run. You get file identity, entropy "
        "and the readable strings inside — summaries of a binary are indicative only."
    )


def _entropy(data: bytes) -> float:
    if not data:
        return 0.0
    import math
    from collections import Counter

    counts = Counter(data)
    total = len(data)
    return -sum((c / total) * math.log2(c / total) for c in counts.values())


def _split_into_pages(text: str, words_per_page: int = 350) -> List[str]:
    """Group text into page-sized blocks so page charts work for any format."""
    if not text:
        return []
    words = text.split()
    return [
        " ".join(words[i : i + words_per_page])
        for i in range(0, len(words), words_per_page)
    ]


# --------------------------------------------------------------------------- #
# Entry point
# --------------------------------------------------------------------------- #

def extract_document(data: bytes, filename: str) -> Document:
    extension = os.path.splitext(filename)[1].lower().lstrip(".")
    doc = Document(
        name=filename,
        extension=extension,
        size_bytes=len(data),
        sha256=_sha256(data),
    )

    if extension == "pdf":
        _read_pdf(data, doc)
    elif extension == "docx":
        _read_docx(data, doc)
    elif extension == "doc":
        _read_legacy_doc(data, doc)
    elif extension in TEXT_EXTENSIONS:
        _read_plain(data, doc)
    elif extension in IMAGE_EXTENSIONS:
        _read_image(data, doc)
    elif extension in BINARY_EXTENSIONS:
        _read_binary(data, doc)
    else:
        doc.notes.append(f".{extension} files are not supported yet.")

    return doc
